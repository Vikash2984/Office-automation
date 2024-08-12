import pandas as pd
from docx import Document
import os

def read_student_data(excel_file):
    """
    Reads student data from an Excel file and returns a list of dictionaries.
    Converts all data to strings where necessary.
    """
    student_data = []
    if os.path.isfile(excel_file):  # Check if file exists
        df = pd.read_excel(excel_file, engine='openpyxl')
        # Convert all data to string
        student_data = df.astype(str).to_dict(orient='records')
    else:
        print(f"File not found: {excel_file}")
    return student_data

def generate_certificate(template_path, output_folder, student_data):
    """
    Generates certificates for students using the given Word template and student data.
    """
    for student in student_data:
        # Load the Word template
        doc = Document(template_path)

        # Replace placeholders with student data
        for paragraph in doc.paragraphs:
            if '{name}' in paragraph.text:
                paragraph.text = paragraph.text.replace('{name}', student['name'])
            if '{course}' in paragraph.text:
                paragraph.text = paragraph.text.replace('{course}', student['course'])
            if '{date}' in paragraph.text:
                paragraph.text = paragraph.text.replace('{date}', student['date'])

        # Save the certificate with student's name
        output_path = os.path.join(output_folder, f"{student['name']}_certificate.docx")
        doc.save(output_path)
        print(f"Generated certificate for {student['name']}")

# Paths
excel_file = r'D:\vsCode\Python\Auto\certificate-generator\Generator-word\XLSS-Generator\data.xlsx'
template_path = r'D:\vsCode\Python\Auto\certificate-generator\Generator-word\template.docx'
output_folder = r'D:\vsCode\Python\Auto\certificate-generator\docx'

# Read student data from Excel
student_data = read_student_data(excel_file)

# Generate certificates
generate_certificate(template_path, output_folder, student_data)
