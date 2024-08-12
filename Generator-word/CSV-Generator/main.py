import csv
from docx import Document
import os

def read_student_data(csv_file):
    """
    Reads student data from a CSV file and returns a list of dictionaries.
    """
    student_data = []
    if os.path.isfile(csv_file):  # Check if file exists
        with open(csv_file, mode='r', newline='', encoding='utf-8') as file:
            reader = csv.DictReader(file)
            for row in reader:
                student_data.append(row)
    else:
        print(f"File not found: {csv_file}")
    return student_data

def generate_certificate(template_path, output_folder, student_data):
    """
    Generates certificates for students using the given Word template and student data.
    """
    for student in student_data:
        # Load the Word template
        doc = Document(template_path)

        # Replace placeholders with student data
        for paragraph in doc.paragraphs:
            if '{name}' in paragraph.text:
                paragraph.text = paragraph.text.replace('{name}', student['name'])
            if '{course}' in paragraph.text:
                paragraph.text = paragraph.text.replace('{course}', student['course'])
            if '{date}' in paragraph.text:
                paragraph.text = paragraph.text.replace('{date}', student['date'])

        # Save the certificate with student's name
        output_path = os.path.join(output_folder, f"{student['name']}_certificate.docx")
        doc.save(output_path)
        print(f"Generated certificate for {student['name']}")

# Paths
csv_file = r'D:\vsCode\Python\Auto\certificate-generator\Generator-word\CSV-Generator\data.csv'
template_path = r'D:\vsCode\Python\Auto\certificate-generator\Generator-word\template.docx'
output_folder = r'D:\vsCode\Python\Auto\certificate-generator\docx'

# Read student data from CSV
student_data = read_student_data(csv_file)

# Generate certificates
generate_certificate(template_path, output_folder, student_data)
